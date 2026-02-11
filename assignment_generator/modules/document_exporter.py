"""
Document Exporter - Generates polished DOCX and PDF files from assignment content.
"""

import os
import logging
from datetime import datetime

logger = logging.getLogger(__name__)


class DocumentExporter:
    """Exports generated assignments to DOCX and PDF formats."""

    def __init__(self, output_dir="outputs"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def export_docx(self, content, filename=None):
        """
        Export assignment content to a formatted DOCX file.

        content: dict with keys 'title', 'sections' (list of {heading, content}), 'bibliography'
        """
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.style import WD_STYLE_TYPE

        doc = Document()

        # --- Page setup ---
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # --- Styles ---
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)
        paragraph_format = style.paragraph_format
        paragraph_format.space_after = Pt(0)
        paragraph_format.space_before = Pt(0)
        paragraph_format.line_spacing = 2.0  # Double spacing

        # --- Title ---
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_para.space_after = Pt(24)
        title_run = title_para.add_run(content["title"])
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_run.font.name = 'Times New Roman'

        # --- Date ---
        date_para = doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_para.space_after = Pt(24)
        date_run = date_para.add_run(datetime.now().strftime("%B %d, %Y"))
        date_run.font.size = Pt(12)
        date_run.font.name = 'Times New Roman'

        # --- Sections ---
        for section_data in content["sections"]:
            # Section heading
            heading_para = doc.add_paragraph()
            heading_para.space_before = Pt(18)
            heading_para.space_after = Pt(12)
            heading_run = heading_para.add_run(section_data["heading"])
            heading_run.bold = True
            heading_run.font.size = Pt(14)
            heading_run.font.name = 'Times New Roman'

            # Section content - split into paragraphs
            paragraphs = section_data["content"].split("\n\n")
            for para_text in paragraphs:
                para_text = para_text.strip()
                if not para_text:
                    continue
                para = doc.add_paragraph()
                para.paragraph_format.first_line_indent = Inches(0.5)
                para.paragraph_format.space_after = Pt(6)
                run = para.add_run(para_text)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)

        # --- Bibliography ---
        if content.get("bibliography"):
            doc.add_page_break()
            bib_heading = doc.add_paragraph()
            bib_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            bib_heading.space_after = Pt(18)
            bib_run = bib_heading.add_run("References")
            bib_run.bold = True
            bib_run.font.size = Pt(14)
            bib_run.font.name = 'Times New Roman'

            for ref in content["bibliography"]:
                ref_para = doc.add_paragraph()
                ref_para.paragraph_format.left_indent = Inches(0.5)
                ref_para.paragraph_format.first_line_indent = Inches(-0.5)
                ref_para.paragraph_format.space_after = Pt(6)
                ref_run = ref_para.add_run(ref)
                ref_run.font.name = 'Times New Roman'
                ref_run.font.size = Pt(12)

        # --- Save ---
        if not filename:
            safe_title = "".join(c if c.isalnum() or c in " -_" else "" for c in content["title"])[:50]
            filename = f"{safe_title}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"

        filepath = os.path.join(self.output_dir, filename)
        doc.save(filepath)
        logger.info(f"DOCX exported: {filepath}")
        return filepath

    def export_pdf(self, content, filename=None):
        """
        Export assignment content to a formatted PDF file.
        """
        from fpdf import FPDF

        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=25)
        pdf.add_page()
        pdf.set_margins(25, 25, 25)

        # --- Title ---
        pdf.set_font("Times", "B", 16)
        pdf.cell(0, 12, content["title"], ln=True, align="C")
        pdf.ln(4)

        # --- Date ---
        pdf.set_font("Times", "", 12)
        pdf.cell(0, 10, datetime.now().strftime("%B %d, %Y"), ln=True, align="C")
        pdf.ln(8)

        # --- Sections ---
        for section_data in content["sections"]:
            # Heading
            pdf.set_font("Times", "B", 14)
            pdf.cell(0, 10, section_data["heading"], ln=True)
            pdf.ln(4)

            # Content
            pdf.set_font("Times", "", 12)
            paragraphs = section_data["content"].split("\n\n")
            for para_text in paragraphs:
                para_text = para_text.strip()
                if not para_text:
                    continue
                # Encode to latin-1 safe text for fpdf
                safe_text = para_text.encode('latin-1', 'replace').decode('latin-1')
                pdf.multi_cell(0, 7, "     " + safe_text)
                pdf.ln(3)

        # --- Bibliography ---
        if content.get("bibliography"):
            pdf.add_page()
            pdf.set_font("Times", "B", 14)
            pdf.cell(0, 12, "References", ln=True, align="C")
            pdf.ln(6)

            pdf.set_font("Times", "", 12)
            for ref in content["bibliography"]:
                safe_ref = ref.encode('latin-1', 'replace').decode('latin-1')
                pdf.multi_cell(0, 7, safe_ref)
                pdf.ln(3)

        # --- Save ---
        if not filename:
            safe_title = "".join(c if c.isalnum() or c in " -_" else "" for c in content["title"])[:50]
            filename = f"{safe_title}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"

        filepath = os.path.join(self.output_dir, filename)
        pdf.output(filepath)
        logger.info(f"PDF exported: {filepath}")
        return filepath

    def get_word_count(self, content):
        """Count total words in the generated content."""
        total = 0
        for section in content.get("sections", []):
            total += len(section.get("content", "").split())
        return total
